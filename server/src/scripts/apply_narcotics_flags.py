import os
import pymongo
from openpyxl import load_workbook

def load_env_mongodb_uri():
    env_path = r"D:\Projects\Medikart\server\.env"
    if os.path.exists(env_path):
        with open(env_path, "r", encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if line.startswith("MONGODB_URI=") and not line.startswith("MONGODB_URI_TEST="):
                    return line.split("=", 1)[1].strip()
    return os.environ.get("MONGODB_URI", "mongodb://127.0.0.1:27017/medikart_dev")

def main():
    excel_path = r"D:\Projects\Medikart\medikart_potential_narcotics_review_list.xlsx"
    if not os.path.exists(excel_path):
        print(f"[-] Excel file not found at {excel_path}")
        return

    print("=" * 80)
    print("[MEDIKART] APPLYING CONFIRMED NARCOTIC FLAGS TO DATABASE")
    print("=" * 80)

    wb = load_workbook(excel_path)
    ws = wb.active

    mongo_uri = load_env_mongodb_uri()
    client = pymongo.MongoClient(mongo_uri)
    try:
        db = client.get_default_database()
        if db is None:
            db = client["medikart_dev"]
    except Exception:
        db = client["medikart_dev"]

    products_col = db["products"]

    updated_count = 0
    skipped_count = 0

    # Header is row 1. Data starts at row 2.
    # Col 2: Name, Col 4: SKU, Col 11: Client Confirmation
    for row in range(2, ws.max_row + 1):
        name = ws.cell(row=row, column=2).value
        sku = ws.cell(row=row, column=4).value
        confirmation = str(ws.cell(row=row, column=11).value or "").strip().upper()

        if confirmation in ("YES", "Y", "TRUE", "1"):
            # Update product in MongoDB
            result = products_col.update_one(
                {"sku": str(sku).strip()},
                {"$set": {"isNarcotic": True, "requiresPrescription": True}}
            )
            if result.matched_count > 0:
                print(f"[+] Flagged as Narcotic: {name} (SKU: {sku})")
                updated_count += 1
            else:
                print(f"[!] Warning: SKU {sku} ({name}) not found in database")
        else:
            skipped_count += 1

    print("\n" + "=" * 80)
    print(f"[*] Summary: {updated_count} products flagged as Narcotic (isNarcotic=true).")
    print(f"[*] Skipped: {skipped_count} products.")
    print("=" * 80)

if __name__ == "__main__":
    main()
